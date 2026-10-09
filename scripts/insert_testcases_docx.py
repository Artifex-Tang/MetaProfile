# -*- coding: utf-8 -*-
"""Insert 7 page-interaction test case tables under 4.6.3 (before 4.6.4 heading)
into 第三方测试大纲0712-课题四标记.docx. 8-row 2-col format matching TC4.2-01."""
import sys, io, copy
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.path.insert(0, 'scripts')
from testcases_draft import CASES  # id,name,pre,content,method,tool,expect,terminate

SRC = '第三方测试大纲0712-课题四标记.docx'
OUT = '第三方测试大纲0712-课题四标记_含页面交互用例.docx'
doc = docx.Document(SRC)

# 1. find 4.6.4 body heading (Heading 3, text startswith 4.6.4) — skip TOC entries
ref = None
for p in doc.paragraphs:
    if p.style and p.style.name == 'Heading 3' and p.text.strip().startswith('4.6.4'):
        ref = p
        break
assert ref is not None, "4.6.4 heading not found"
ref_el = ref._element
print(f"insert before: {ref.text.strip()[:30]}")

# helper: create element, move before ref
def insert_before(el):
    ref_el.addprevious(el)

# 2. new heading 4.6.3.4
h = doc.add_paragraph('4.6.3.4 工具页面交互功能验证（页面交互测试用例）', style='Heading 4')
insert_before(h._element)
# intro line
intro = doc.add_paragraph(
    '本节针对产业技术识别模型及工具集（课题四）部署系统（MetaProfile，http://47.120.41.179）'
    '的7个工具页面交互功能进行验证，测试用例覆盖系统Swagger接口（http://47.120.41.179/docs）'
    '对应的页面交互场景，采用黑盒测试方法。')
intro.paragraph_format.space_after = Pt(6)
for r in intro.runs:
    r.font.size = Pt(10.5)
insert_before(intro._element)

# 3. field order (TC4.2-01 format, 8 rows)
ROWS = [
    ('用例编号', 'id'),
    ('用例名称', 'name'),
    ('前置条件', 'pre'),
    ('测试内容及评判标准', 'content'),
    ('测试方法', 'method'),
    ('依赖的测试工具', 'tool'),
    ('预期结果', 'expect'),
    ('测试项终止条件', 'terminate'),
]

def set_cell_font(cell, bold=False, size=10.5):
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = '宋体'
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_table_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000')
        borders.append(e)
    tblPr.append(borders)

for case in CASES:
    tbl = doc.add_table(rows=len(ROWS), cols=2)
    set_table_borders(tbl)
    tbl.autofit = False
    # column widths: label ~3.2cm, value ~13cm
    for r_i, (label, key) in enumerate(ROWS):
        c0 = tbl.rows[r_i].cells[0]
        c1 = tbl.rows[r_i].cells[1]
        c0.width = Cm(3.2); c1.width = Cm(13.0)
        # label
        c0.text = ''
        p0 = c0.paragraphs[0]
        run0 = p0.add_run(label)
        set_cell_font(c0, bold=True)
        shade_cell(c0, 'F2F2F2')
        # value (support multi-line via \n)
        c1.text = ''
        val = str(case.get(key, ''))
        lines = val.split('\n')
        for li, line in enumerate(lines):
            p1 = c1.paragraphs[0] if li == 0 else c1.add_paragraph()
            p1.add_run(line)
            set_cell_font(c1, bold=False)
    # move table before ref
    insert_before(tbl._element)
    # spacer paragraph after each table
    sp = doc.add_paragraph('')
    sp.paragraph_format.space_after = Pt(4)
    insert_before(sp._element)
    print(f"inserted {case['id']} {case['name']}")

doc.save(OUT)
print(f"=== saved {OUT}, tables={len(CASES)} ===")
